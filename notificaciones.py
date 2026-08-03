"""
Notificaciones del monitoreo de adjudicadas: correo (SMTP) y export a Excel.

El envío de correo usa SMTP estándar (compatible con Gmail vía contraseña de
aplicación). Se configura con variables de entorno para que funcione tanto en
local como en la automatización de GitHub Actions:

    SMTP_HOST       (por defecto smtp.gmail.com)
    SMTP_PORT       (por defecto 587, STARTTLS)
    SMTP_USER       usuario/casilla remitente
    SMTP_PASSWORD   contraseña de aplicación (NO la contraseña normal de Gmail)
    EMAIL_FROM      remitente (por defecto = SMTP_USER)
    EMAIL_TO        destinatario(s), separados por coma

Para Gmail: activar verificación en dos pasos y crear una "Contraseña de
aplicación" en https://myaccount.google.com/apppasswords
"""

from __future__ import annotations

import os
import smtplib
from datetime import datetime
from email.message import EmailMessage
from io import BytesIO
from typing import Any


# ---------------------------------------------------------------------------
# Utilidades de formato
# ---------------------------------------------------------------------------

def formato_clp(monto: Any) -> str:
    if not monto:
        return "—"
    try:
        return f"${float(monto):,.0f}".replace(",", ".")
    except (ValueError, TypeError):
        return "—"


def _proveedores_texto(registro: dict[str, Any]) -> str:
    adj = registro.get("Adjudicatarios") or []
    nombres = [a.get("proveedor", "") for a in adj if a.get("proveedor")]
    return " · ".join(nombres) if nombres else "No informado"


# ---------------------------------------------------------------------------
# Correo
# ---------------------------------------------------------------------------

class NotificacionError(Exception):
    """Error al enviar una notificación."""


def config_smtp_desde_entorno() -> dict[str, str]:
    return {
        "host": os.environ.get("SMTP_HOST", "smtp.gmail.com"),
        "port": os.environ.get("SMTP_PORT", "587"),
        "user": os.environ.get("SMTP_USER", ""),
        "password": os.environ.get("SMTP_PASSWORD", ""),
        "from": os.environ.get("EMAIL_FROM", os.environ.get("SMTP_USER", "")),
        "to": os.environ.get("EMAIL_TO", ""),
    }


def smtp_configurado(cfg: dict[str, str] | None = None) -> bool:
    cfg = cfg or config_smtp_desde_entorno()
    return bool(cfg.get("user") and cfg.get("password") and cfg.get("to"))


def _html_resumen(registros: list[dict[str, Any]]) -> str:
    filas = ""
    for r in registros:
        filas += f"""
        <tr>
            <td style="padding:8px;border-bottom:1px solid #eee;">
                <b>{r.get('Nombre', '')}</b><br>
                <span style="color:#7A7A7A;font-size:12px;">
                    {r.get('NombreOrganismo', '')} · {r.get('Region', '')} ·
                    <code>{r.get('CodigoExterno', '')}</code>
                </span>
            </td>
            <td style="padding:8px;border-bottom:1px solid #eee;">
                {_proveedores_texto(r)}
            </td>
            <td style="padding:8px;border-bottom:1px solid #eee;text-align:right;white-space:nowrap;">
                <b>{formato_clp(r.get('MontoAdjudicado'))}</b>
            </td>
            <td style="padding:8px;border-bottom:1px solid #eee;text-align:center;">
                <a href="{r.get('UrlMercadoPublico', '#')}"
                   style="color:#A6212B;text-decoration:none;">Ver ↗</a>
            </td>
        </tr>"""

    total = sum(float(r.get("MontoAdjudicado") or 0) for r in registros)

    return f"""
    <div style="font-family:Arial,Helvetica,sans-serif;color:#1A1A1A;max-width:760px;">
        <div style="background:#A6212B;color:white;padding:16px 20px;border-radius:6px 6px 0 0;">
            <div style="font-size:18px;font-weight:900;">RCH · Monitoreo de Adjudicaciones</div>
            <div style="font-size:13px;opacity:.9;">
                {len(registros)} nueva(s) adjudicación(es) en tu rubro ·
                {datetime.now().strftime('%d-%m-%Y %H:%M')}
            </div>
        </div>
        <table style="border-collapse:collapse;width:100%;font-size:14px;
                      border:1px solid #eee;border-top:none;">
            <thead>
                <tr style="background:#EDEAE5;text-align:left;">
                    <th style="padding:8px;">Obra / Licitación</th>
                    <th style="padding:8px;">Adjudicatario(s)</th>
                    <th style="padding:8px;text-align:right;">Monto</th>
                    <th style="padding:8px;text-align:center;">MP</th>
                </tr>
            </thead>
            <tbody>{filas}</tbody>
            <tfoot>
                <tr style="background:#fafafa;">
                    <td colspan="2" style="padding:10px;font-weight:700;">Monto total adjudicado</td>
                    <td style="padding:10px;text-align:right;font-weight:900;color:#A6212B;">
                        {formato_clp(total)}
                    </td>
                    <td></td>
                </tr>
            </tfoot>
        </table>
        <p style="color:#7A7A7A;font-size:12px;margin-top:14px;">
            Sistema de Análisis de Licitaciones · RCH · Aviso automático de Mercado Público.
        </p>
    </div>"""


def enviar_email_resumen(
    registros: list[dict[str, Any]],
    cfg: dict[str, str] | None = None,
) -> None:
    """Envía un correo HTML con las adjudicaciones nuevas.

    Lanza NotificacionError si la configuración SMTP está incompleta o el
    envío falla.
    """
    if not registros:
        return

    cfg = cfg or config_smtp_desde_entorno()
    if not smtp_configurado(cfg):
        raise NotificacionError(
            "Configuración SMTP incompleta (faltan SMTP_USER, SMTP_PASSWORD o EMAIL_TO)."
        )

    mensaje = EmailMessage()
    mensaje["Subject"] = (
        f"[RCH] {len(registros)} nueva(s) adjudicación(es) en tu rubro"
    )
    mensaje["From"] = cfg["from"]
    destinatarios = [d.strip() for d in cfg["to"].split(",") if d.strip()]
    mensaje["To"] = ", ".join(destinatarios)

    # Cuerpo texto plano (respaldo) + HTML
    lineas = [
        f"- {r.get('Nombre', '')} | {_proveedores_texto(r)} | "
        f"{formato_clp(r.get('MontoAdjudicado'))} | {r.get('UrlMercadoPublico', '')}"
        for r in registros
    ]
    mensaje.set_content(
        "Nuevas adjudicaciones detectadas en tu rubro:\n\n" + "\n".join(lineas)
    )
    mensaje.add_alternative(_html_resumen(registros), subtype="html")

    try:
        puerto = int(cfg.get("port", "587"))
        with smtplib.SMTP(cfg["host"], puerto, timeout=30) as servidor:
            servidor.starttls()
            servidor.login(cfg["user"], cfg["password"])
            servidor.send_message(mensaje)
    except (smtplib.SMTPException, OSError) as exc:
        raise NotificacionError(f"Fallo al enviar el correo: {exc}") from exc


# ---------------------------------------------------------------------------
# Export a Excel
# ---------------------------------------------------------------------------

def generar_excel_adjudicadas(registros: list[dict[str, Any]]) -> bytes:
    """Genera un Excel con formato RCH y devuelve los bytes del archivo."""
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    COLS = [
        "Código", "Obra / Licitación", "Organismo", "Región",
        "Rubro (ONU)", "Adjudicatario(s)", "Monto adjudicado (CLP)",
        "Fecha adjudicación", "URL Mercado Público",
    ]
    WIDTHS = [16, 42, 30, 24, 12, 34, 22, 18, 40]

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Adjudicadas"
    ws.sheet_view.showGridLines = False

    header_fill = PatternFill("solid", fgColor="1A1A1A")
    accent_fill = PatternFill("solid", fgColor="A6212B")
    band_fill = PatternFill("solid", fgColor="EDEAE5")
    header_font = Font(name="Arial", bold=True, color="FFFFFF", size=11)
    title_font = Font(name="Arial", bold=True, color="FFFFFF", size=15)
    cell_font = Font(name="Arial", size=10)
    thin = Side(style="thin", color="CCCCCC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    ws.merge_cells("A1:I1")
    t = ws.cell(row=1, column=1, value="RCH · Obras adjudicadas monitoreadas")
    t.font = title_font
    t.fill = accent_fill
    t.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[1].height = 30

    ws.merge_cells("A2:I2")
    s = ws.cell(
        row=2, column=1,
        value=f"Generado el {datetime.now().strftime('%d-%m-%Y %H:%M')} · {len(registros)} adjudicaciones",
    )
    s.font = Font(name="Arial", italic=True, color="1A1A1A", size=10)
    s.fill = band_fill
    s.alignment = Alignment(horizontal="left", vertical="center", indent=1)

    fila_header = 4
    for c, name in enumerate(COLS, start=1):
        cell = ws.cell(row=fila_header, column=c, value=name)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        cell.border = border
    ws.row_dimensions[fila_header].height = 26

    r = fila_header + 1
    for i, reg in enumerate(registros):
        adj = reg.get("Adjudicatarios") or []
        proveedores = " · ".join(a.get("proveedor", "") for a in adj if a.get("proveedor"))
        valores = [
            reg.get("CodigoExterno", ""),
            reg.get("Nombre", ""),
            reg.get("NombreOrganismo", ""),
            reg.get("Region", ""),
            str(reg.get("CodigoProductoONU", ""))[:2],
            proveedores or "No informado",
            reg.get("MontoAdjudicado") or "",
            str(reg.get("FechaAdjudicacion", ""))[:10],
            reg.get("UrlMercadoPublico", ""),
        ]
        for c, val in enumerate(valores, start=1):
            cell = ws.cell(row=r, column=c, value=val)
            cell.font = cell_font
            cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
            cell.border = border
            if i % 2 == 1:
                cell.fill = band_fill
            if c == 7 and isinstance(val, (int, float)):
                cell.number_format = "#,##0"
        r += 1

    for c, w in enumerate(WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(c)].width = w
    if registros:
        ws.auto_filter.ref = f"A{fila_header}:{get_column_letter(len(COLS))}{r-1}"
    ws.freeze_panes = ws.cell(row=fila_header + 1, column=1)

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
